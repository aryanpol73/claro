import os
from pathlib import Path
from typing import List
from langchain_core.documents import Document


def load_pptx(file_path: str) -> List[Document]:
    """
    Extract text and metadata from PowerPoint (.pptx) files slide-by-slide.
    Iterates shapes, text frames, tables, and speaker notes per slide.
    Maps each slide to a Document with 1-indexed page/slide number.
    """
    from pptx import Presentation

    documents: List[Document] = []
    prs = Presentation(file_path)
    total_slides = len(prs.slides)
    source_name = os.path.basename(file_path)

    for idx, slide in enumerate(prs.slides):
        slide_number = idx + 1
        text_parts: List[str] = []

        # Extract title shape if present
        if slide.shapes.title and slide.shapes.title.text.strip():
            text_parts.append(f"## {slide.shapes.title.text.strip()}")

        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        text_parts.append(line)
            elif shape.has_table:
                table_lines = []
                for row in shape.table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_lines.append(" | ".join(row_cells))
                if table_lines:
                    text_parts.append("\n".join(table_lines))

        # Check speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                text_parts.append(f"[Speaker Notes: {notes}]")

        slide_content = "\n\n".join(text_parts).strip()
        if not slide_content:
            # Keep slide with minimal placeholder so citations / page count align
            slide_content = f"Slide {slide_number}"

        doc = Document(
            page_content=slide_content,
            metadata={
                "source": source_name,
                "page": slide_number,
                "total_pages": total_slides,
            },
        )
        documents.append(doc)

    return documents
