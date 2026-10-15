import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

# Ensure UTF-8 output
sys.stdout.reconfigure(encoding="utf-8")

# Step 1: Create a sample PPTX file
sample_pptx = Path("backend/uploads/sample_quantum.pptx")
prs = Presentation()

# Slide 1: Title
slide1 = prs.slides.add_slide(prs.slide_layouts[0])
slide1.shapes.title.text = "Introduction to Quantum Computing"
slide1.placeholders[1].text = "A Primer on Qubits and Superposition"

# Slide 2: Qubits & Superposition
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
slide2.shapes.title.text = "Qubits and Superposition"
tf2 = slide2.shapes.placeholders[1].text_frame
tf2.text = "A qubit is the fundamental unit of quantum information."
p = tf2.add_paragraph()
p.text = "Unlike classical bits that are strictly 0 or 1, a qubit can exist in a linear combination of both states simultaneously, known as superposition."
p2 = tf2.add_paragraph()
p2.text = "The state of a single qubit is mathematically represented as |psi> = alpha|0> + beta|1>, where |alpha|^2 + |beta|^2 = 1."

# Slide 3: Quantum Gates
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
slide3.shapes.title.text = "Quantum Gates"
tf3 = slide3.shapes.placeholders[1].text_frame
tf3.text = "Quantum logic gates are reversible transformations performed on qubits."
p3 = tf3.add_paragraph()
p3.text = "The Hadamard gate (H gate) creates an equal superposition from a basis state."
p4 = tf3.add_paragraph()
p4.text = "The CNOT gate (Controlled-NOT) acts on two qubits to create quantum entanglement."

prs.save(str(sample_pptx))
print(f"[OK] Created sample PPTX with {len(prs.slides)} slides at {sample_pptx}")
